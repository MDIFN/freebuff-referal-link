import json
import re
import unicodedata
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCES = ROOT / "sources"


def clean(value):
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value))
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", text)
    return re.sub(r"\s+", " ", text).strip()


def key(*values):
    return tuple(clean(value).casefold() for value in values)


def workbook_records():
    workbook = load_workbook(
        SOURCES / "Hyderabad_Telangana_Medicine_Master_Phase3.xlsx",
        read_only=True,
        data_only=True,
    )
    sheet = workbook.active
    records = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        name, strength, indication = (clean(row[index]) for index in (1, 2, 3))
        if not name:
            continue
        record_key = key(name, strength, indication)
        records.setdefault(
            record_key,
            {
                "type": "approval",
                "name": name,
                "strength": strength,
                "indication": indication,
                "source": "CDSCO approval master",
            },
        )
    workbook.close()
    return list(records.values())


def tender_records():
    records = {}
    layout = None
    with pdfplumber.open(SOURCES / "drugs names in PDF.pdf") as document:
        for page in document.pages:
            for table in page.extract_tables() or []:
                header = clean(" ".join(cell or "" for cell in table[0])) if table else ""
                if "Generic Name of Drug" in header:
                    layout = "drug"
                elif "Generic Name of Item" in header:
                    layout = "item"
                if layout is None:
                    continue

                for row in table:
                    if not row or len(row) < 5 or not clean(row[0]).isdigit():
                        continue
                    name, specification, unit_size = (clean(row[index]) for index in (2, 3, 4))
                    if not name:
                        continue
                    record_key = key(name, specification, unit_size)
                    record = records.setdefault(
                        record_key,
                        {
                            "type": "tender",
                            "name": name,
                            "genericName": name,
                            "strength": "",
                            "specification": specification,
                            "indication": "",
                            "unitSize": unit_size,
                            "packSize": "",
                            "packingType": "",
                            "packingStandard": "",
                            "source": "PMBI drug tender",
                        },
                    )
                    if layout == "drug" and len(row) > 5:
                        record["packSize"] = clean(row[5])
                    elif layout == "item" and len(row) > 6:
                        record["packingType"] = clean(row[5])
                        record["packingStandard"] = clean(row[6])
    return list(records.values())


def main():
    records = workbook_records() + tender_records()
    if not records:
        raise SystemExit("No medicine records extracted; check source files and layouts.")

    payload = json.dumps(records, ensure_ascii=True, separators=(",", ":"))
    output = ROOT / "drug-reference.js"
    output.write_text(f"window.DRUG_REFERENCE = {payload};\n", encoding="utf-8")
    print(f"Wrote {len(records)} reference entries to {output.name} ({output.stat().st_size:,} bytes).")


if __name__ == "__main__":
    main()